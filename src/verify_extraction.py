"""Compare extracted MAA figures with the hand-built answer key workbook."""
from __future__ import annotations

import json
from pathlib import Path

# Workbook rows (MAA Reconciliation tab) and the matching extraction field or category.
TOTAL_ROWS = {"Net income (start)": 8, "FFO": 15, "Core FFO": 23, "Core AFFO": 27, "Diluted shares + units": 30}
CATEGORY_ROWS = {
    "re_depreciation": 9, "property_sale_gain_loss": 10, "jv_depreciation": 11, "noncontrolling_interest": 13,
    "derivative_mtm": 18, "investment_gain_loss": 19, "casualty": 20, "legal": 21, "recurring_capex": 26,
}
COLUMNS = {"2025Q4": 2, "2026Q1": 3, "2026Q2": 4}


def compare_maa(workbook: Path, extract_dir: Path) -> int:
    from openpyxl import load_workbook

    ws = load_workbook(workbook, data_only=True)["MAA Reconciliation"]
    mismatches = checked = 0
    for quarter, col in COLUMNS.items():
        path = Path(extract_dir) / f"MAA_{quarter}.json"
        if not path.exists():
            print(f"  {quarter}: no extraction yet")
            continue
        ext = json.loads(path.read_text())["extraction"]
        lines = ext["ffo_adjustments"] + ext["core_adjustments"] + ext["affo_adjustments"]
        got = {
            "Net income (start)": ext["starting_line"]["amount"], "FFO": ext["ffo_total"]["amount"],
            "Core FFO": (ext.get("core_total") or {}).get("amount"), "Core AFFO": (ext.get("affo_total") or {}).get("amount"),
            "Diluted shares + units": ext["diluted_shares"]["amount"],
        }
        for cat in CATEGORY_ROWS:
            got[cat] = sum(l["amount"] for l in lines if l["category"] == cat)
        bad = []
        for name, row in {**TOTAL_ROWS, **CATEGORY_ROWS}.items():
            expected = ws.cell(row, col).value or 0
            checked += 1
            if got.get(name) is None or abs((got[name] or 0) - expected) > 1:
                bad.append(f"{name}: expected {expected:,.0f}, extracted {got.get(name)}")
        mismatches += len(bad)
        print(f"  {quarter}: " + ("all 14 figures match" if not bad else "MISMATCH\n    - " + "\n    - ".join(bad)))
    print(f"  Overall: {checked - mismatches}/{checked} figures match the answer key")
    return mismatches
